import os
import email
from email.policy import default
from PyPDF2 import PdfReader
import docx
import openpyxl
from pinecone import Pinecone, ServerlessSpec
from openai import OpenAI

# Hardcoded API keys
PC_API_KEY = "pcsk_5CwF7M_2c71gSS2ogeQnpVRd3TJWHrj69hJBVGn1uZxNqtbSkgeXXszh6Q2dDpvghvA6sF"  # Replace with your Pinecone API key
PC_ENVIRONMENT = "us-east-1"  # Replace with your Pinecone environment
OPENAI_API_KEY = "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"  # Replace with your OpenAI API key

# Initialize OpenAI
client = OpenAI(api_key=OPENAI_API_KEY)

# Initialize Pinecone
pc = Pinecone(api_key=PC_API_KEY, environment=PC_ENVIRONMENT)

# Create or connect to a Pinecone index
index_name = "boyscout-gpt-t125"
# if index_name not in pc.list_indexes():
#     pc.create_index(
#         index_name,
#         dimension=1536,
#         spec=ServerlessSpec(cloud="aws", region=PC_ENVIRONMENT),
#     )  # OpenAI embedding dimension is 1536
index = pc.Index(index_name)

# Directories
uncleaned_dir = "Uncleaned_Emails"
cleaned_dir = "Cleaned_Emails"
os.makedirs(cleaned_dir, exist_ok=True)

# Allowed document types
allowed_extensions = {".pdf", ".docx", ".xlsx"}


# Function to extract text from PDFs
def extract_pdf_content(file_path):
    content = []
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            content.append(page.extract_text())
    except Exception as e:
        content.append(f"Error reading PDF: {e}")
    return "\n".join(content)


# Function to extract text from Word documents
def extract_docx_content(file_path):
    content = []
    try:
        doc = docx.Document(file_path)
        for paragraph in doc.paragraphs:
            content.append(paragraph.text)
    except Exception as e:
        content.append(f"Error reading DOCX: {e}")
    return "\n".join(content)


# Function to extract text from Excel files
def extract_xlsx_content(file_path):
    content = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb:
            content.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                content.append("\t".join(str(cell or "") for cell in row))
    except Exception as e:
        content.append(f"Error reading XLSX: {e}")
    return "\n".join(content)


# Function to remove empty lines
def remove_empty_lines(text):
    return "\n".join(line for line in text.split("\n") if line.strip() != "")


# Function to clean an email
def clean_email(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        msg = email.message_from_file(file, policy=default)

    # Extract metadata
    subject = msg.get("Subject", "No Subject")
    sender = msg.get("From", "Unknown Sender")
    date = msg.get("Date", "No Date")
    body = ""
    attachments_content = []

    # Extract email body and attachments
    if msg.is_multipart():
        for part in msg.walk():
            content_type = part.get_content_type()
            disposition = part.get("Content-Disposition", "")

            if content_type == "text/plain" and "attachment" not in disposition:
                body = part.get_payload(decode=True).decode(
                    part.get_content_charset() or "utf-8"
                )
            elif "attachment" in disposition:
                filename = part.get_filename()
                if filename:
                    ext = os.path.splitext(filename)[1].lower()
                    if ext in allowed_extensions:
                        file_data = part.get_payload(decode=True)
                        temp_path = os.path.join(cleaned_dir, filename)
                        with open(temp_path, "wb") as temp_file:
                            temp_file.write(file_data)
                        if ext == ".pdf":
                            attachments_content.append(
                                f"--- Attachment Content ({filename}) ---\n"
                                + remove_empty_lines(extract_pdf_content(temp_path))
                            )
                        elif ext == ".docx":
                            attachments_content.append(
                                f"--- Attachment Content ({filename}) ---\n"
                                + remove_empty_lines(extract_docx_content(temp_path))
                            )
                        elif ext == ".xlsx":
                            attachments_content.append(
                                f"--- Attachment Content ({filename}) ---\n"
                                + remove_empty_lines(extract_xlsx_content(temp_path))
                            )
                        os.remove(temp_path)
    else:
        body = msg.get_payload(decode=True).decode(msg.get_content_charset() or "utf-8")

    # Combine cleaned email content
    cleaned_content = f"Subject: {subject}\nFrom: {sender}\nDate: {date}\n\n{remove_empty_lines(body.strip())}"
    if attachments_content:
        cleaned_content += "\n\nAttachments:\n" + "\n".join(attachments_content)

    return cleaned_content, subject, sender, date


# Function to generate embeddings
def embed_text(text):
    response = client.embeddings.create(model="text-embedding-ada-002", input=text)
    return response.data[0].embedding


# Process emails and upload to Pinecone
for filename in os.listdir(uncleaned_dir):
    if filename.endswith(".eml"):
        file_path = os.path.join(uncleaned_dir, filename)
        cleaned_content, subject, sender, date = clean_email(file_path)

        # Save cleaned content
        cleaned_file_path = os.path.join(
            cleaned_dir, f"{os.path.splitext(filename)[0]}.txt"
        )
        with open(cleaned_file_path, "w", encoding="utf-8") as cleaned_file:
            cleaned_file.write(cleaned_content)

        # Generate embedding and upload to Pinecone
        try:
            embedding = embed_text(cleaned_content)
            metadata = {"subject": subject, "from": sender, "date": date}
            index.upsert([(filename, embedding, metadata)])
            print(f"Uploaded {filename} with metadata to Pinecone.")
        except Exception as e:
            print(f"Error uploading {filename}: {e}")

print("Emails cleaned and uploaded to Pinecone with metadata.")
