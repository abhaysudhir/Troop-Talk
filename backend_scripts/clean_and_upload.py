import os
import email
import time
import datetime
import tiktoken
from email.policy import default
from PyPDF2 import PdfReader
import docx
import openpyxl
from pinecone import Pinecone
from openai import OpenAI
from dotenv import load_dotenv

load_dotenv()
print("\n" + "="*80)
print(f"SCRIPT STARTED: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print("="*80)

# API keys and configurations
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_ENV = "us-east-1"

print("\nCONFIGURATION:")
print(f"OpenAI API Key: {'✓ Found' if OPENAI_API_KEY else '✗ Missing'}")
print(f"Pinecone API Key: {'✓ Found' if PINECONE_API_KEY else '✗ Missing'}")
print(f"Pinecone Environment: {PINECONE_ENV}")

# Initialize OpenAI and Pinecone
print("\nINITIALIZING CLIENTS...")
try:
    openai_client = OpenAI(api_key=OPENAI_API_KEY)
    print("✓ OpenAI client initialized successfully")
except Exception as e:
    print(f"✗ Error initializing OpenAI client: {e}")
    exit(1)

try:
    pinecone_client = Pinecone(api_key=PINECONE_API_KEY, environment=PINECONE_ENV)
    print("✓ Pinecone client initialized successfully")
except Exception as e:
    print(f"✗ Error initializing Pinecone client: {e}")
    exit(1)

# Set up Pinecone index
INDEX_NAME = "boyscout-gpt-t125"
print(f"\nCONNECTING TO PINECONE INDEX: {INDEX_NAME}")
try:
    index = pinecone_client.Index(INDEX_NAME)
    print(f"✓ Successfully connected to Pinecone index '{INDEX_NAME}'")
except Exception as e:
    print(f"✗ Error connecting to Pinecone index: {e}")
    exit(1)

# Directories for processing
UNCLEANED_DIR = "Uncleaned_Emails" # Change between Uncleaned_Emails or Uncleaned_General_Info
CLEANED_DIR = "Cleaned_Emails" # Change between Cleaned_Emails or Cleaned_General_Info
print(f"\nDIRECTORIES:")
print(f"Input directory: {os.path.abspath(UNCLEANED_DIR)}")
print(f"Output directory: {os.path.abspath(CLEANED_DIR)}")

if not os.path.exists(UNCLEANED_DIR):
    print(f"✗ Input directory '{UNCLEANED_DIR}' does not exist!")
    exit(1)
else:
    print(f"✓ Input directory exists")
    file_count = len([f for f in os.listdir(UNCLEANED_DIR) if os.path.isfile(os.path.join(UNCLEANED_DIR, f))])
    print(f"  Found {file_count} files in input directory")

os.makedirs(CLEANED_DIR, exist_ok=True)
print(f"✓ Output directory ready")

# Allowed file extensions
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md"}
print(f"\nALLOWED EXTENSIONS: {', '.join(ALLOWED_EXTENSIONS)}")


# Helper functions for content extraction
def extract_pdf_content(file_path):
    print(f"  Extracting content from PDF: {os.path.basename(file_path)}")
    start_time = time.time()
    try:
        reader = PdfReader(file_path)
        page_count = len(reader.pages)
        print(f"  PDF has {page_count} pages")
        
        content = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            content.append(page_text)
            print(f"  Extracted page {i+1}/{page_count} ({len(page_text)} characters)")
        
        full_content = "\n".join(content)
        print(f"  ✓ PDF extraction complete - {len(full_content)} total characters in {time.time() - start_time:.2f} seconds")
        return full_content
    except Exception as e:
        print(f"  ✗ Error reading PDF: {e}")
        return f"Error reading PDF: {e}"


def extract_docx_content(file_path):
    print(f"  Extracting content from DOCX: {os.path.basename(file_path)}")
    start_time = time.time()
    try:
        doc = docx.Document(file_path)
        paragraph_count = len(doc.paragraphs)
        print(f"  DOCX has {paragraph_count} paragraphs")
        
        content = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        print(f"  ✓ DOCX extraction complete - {len(content)} total characters in {time.time() - start_time:.2f} seconds")
        return content
    except Exception as e:
        print(f"  ✗ Error reading DOCX: {e}")
        return f"Error reading DOCX: {e}"


def extract_xlsx_content(file_path):
    print(f"  Extracting content from XLSX: {os.path.basename(file_path)}")
    start_time = time.time()
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet_count = len(wb.sheetnames)
        print(f"  XLSX has {sheet_count} sheets: {', '.join(wb.sheetnames)}")
        
        content = []
        for sheet in wb:
            row_count = sheet.max_row
            col_count = sheet.max_column
            print(f"  Processing sheet '{sheet.title}' with {row_count} rows and {col_count} columns")
            
            content.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                content.append("\t".join(str(cell or "") for cell in row))
        
        full_content = "\n".join(content)
        print(f"  ✓ XLSX extraction complete - {len(full_content)} total characters in {time.time() - start_time:.2f} seconds")
        return full_content
    except Exception as e:
        print(f"  ✗ Error reading XLSX: {e}")
        return f"Error reading XLSX: {e}"


def extract_txt_content(file_path):
    print(f"  Extracting content from TXT: {os.path.basename(file_path)}")
    start_time = time.time()
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
        print(f"  ✓ TXT extraction complete - {len(content)} total characters in {time.time() - start_time:.2f} seconds")
        return content
    except Exception as e:
        print(f"  ✗ Error reading TXT: {e}")
        return f"Error reading TXT: {e}"


def extract_markdown_content(file_path):
    print(f"  Extracting content from Markdown: {os.path.basename(file_path)}")
    start_time = time.time()
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
        print(f"  ✓ Markdown extraction complete - {len(content)} total characters in {time.time() - start_time:.2f} seconds")
        return content
    except Exception as e:
        print(f"  ✗ Error reading Markdown: {e}")
        return f"Error reading Markdown: {e}"


def remove_empty_lines(text):
    print(f"  Removing empty lines from text ({len(text)} characters)")
    lines_before = text.count('\n') + 1
    cleaned_text = "\n".join(line for line in text.split("\n") if line.strip() != "")
    lines_after = cleaned_text.count('\n') + 1
    print(f"  ✓ Removed {lines_before - lines_after} empty lines")
    return cleaned_text


# Function to clean email
def clean_email(file_path):
    print(f"  Processing email file: {os.path.basename(file_path)}")
    start_time = time.time()
    
    with open(file_path, "r", encoding="utf-8") as file:
        msg = email.message_from_file(file, policy=default)

    subject = msg.get("Subject", "No Subject")
    sender = msg.get("From", "Unknown Sender")
    date = msg.get("Date", "No Date")
    
    print(f"  Email metadata:")
    print(f"    Subject: {subject}")
    print(f"    From: {sender}")
    print(f"    Date: {date}")
    
    body = []
    attachments = []

    if msg.is_multipart():
        print(f"  Email is multipart")
        part_count = 0
        for part in msg.walk():
            part_count += 1
            content_type = part.get_content_type()
            disposition = part.get("Content-Disposition", None)
            
            print(f"  Processing part {part_count}: {content_type}, disposition: {disposition}")

            if content_type == "text/plain" and (
                not disposition or "attachment" not in disposition
            ):
                try:
                    part_content = part.get_payload(decode=True).decode(
                        part.get_content_charset() or "utf-8"
                    )
                    body.append(part_content)
                    print(f"    ✓ Added text part ({len(part_content)} characters)")
                except Exception as e:
                    error_msg = f"Error decoding body part: {e}"
                    body.append(error_msg)
                    print(f"    ✗ {error_msg}")
            elif disposition and "attachment" in disposition:
                filename = part.get_filename()
                if filename:
                    print(f"    Found attachment: {filename}")
                    ext = os.path.splitext(filename)[1].lower()
                    if ext in ALLOWED_EXTENSIONS:
                        print(f"    Processing attachment with extension {ext}")
                        temp_path = os.path.join(CLEANED_DIR, filename)
                        with open(temp_path, "wb") as temp_file:
                            temp_file.write(part.get_payload(decode=True))
                        
                        attachment_content = ""
                        if ext == ".pdf":
                            attachment_content = extract_pdf_content(temp_path)
                        elif ext == ".docx":
                            attachment_content = extract_docx_content(temp_path)
                        elif ext == ".xlsx":
                            attachment_content = extract_xlsx_content(temp_path)
                        elif ext == ".txt":
                            attachment_content = extract_txt_content(temp_path)
                        elif ext == ".md":
                            attachment_content = extract_markdown_content(temp_path)
                        
                        attachments.append(attachment_content)
                        print(f"    ✓ Processed attachment: {filename} ({len(attachment_content)} characters)")
                        os.remove(temp_path)
                        print(f"    ✓ Removed temporary file: {temp_path}")
                    else:
                        print(f"    ✗ Skipping attachment with unsupported extension: {ext}")
    else:
        print(f"  Email is single-part")
        try:
            part_content = msg.get_payload(decode=True).decode(
                msg.get_content_charset() or "utf-8"
            )
            body.append(part_content)
            print(f"    ✓ Added email body ({len(part_content)} characters)")
        except Exception as e:
            error_msg = f"Error decoding single-part email: {e}"
            body.append(error_msg)
            print(f"    ✗ {error_msg}")

    cleaned_body = remove_empty_lines("\n".join(body))
    print(f"  Email body: {len(cleaned_body)} characters after cleaning")
    
    attachments_text = ""
    if attachments:
        attachments_text = "\n\n".join(attachments)
        print(f"  Attachments: {len(attachments)} with total {len(attachments_text)} characters")
    
    full_content = f"Subject: {subject}\nFrom: {sender}\nDate: {date}\n\n{cleaned_body}"
    if attachments_text:
        full_content += f"\n\n{attachments_text}"
    
    print(f"  ✓ Email processing complete - {len(full_content)} total characters in {time.time() - start_time:.2f} seconds")
    
    return (
        full_content,
        subject,
        sender,
        date,
    )


def chunk_text(text, chunk_size=400):
    """
    Split text into chunks of approximately chunk_size tokens.
    Ensures splits occur at sentence or paragraph boundaries when possible.
    
    Args:
        text (str): The text to split into chunks
        chunk_size (int): Target number of tokens per chunk (default 400)
    
    Returns:
        list: List of text chunks
    """
    print(f"  Chunking text of {len(text)} characters into ~{chunk_size} token chunks")
    start_time = time.time()
    
    # Initialize tokenizer
    enc = tiktoken.encoding_for_model("text-embedding-ada-002")
    
    # First split by double newlines (paragraphs)
    paragraphs = text.split('\n\n')
    chunks = []
    current_chunk = []
    current_token_count = 0
    
    for paragraph in paragraphs:
        # If paragraph is empty, skip it
        if not paragraph.strip():
            continue
            
        # Get token count for this paragraph
        paragraph_tokens = len(enc.encode(paragraph))
        
        # If a single paragraph is too large, split it by sentences
        if paragraph_tokens > chunk_size:
            sentences = [s.strip() for s in paragraph.replace('\n', ' ').split('. ')]
            for sentence in sentences:
                if not sentence:
                    continue
                sentence_tokens = len(enc.encode(sentence))
                
                # If adding this sentence would exceed chunk size, start a new chunk
                if current_token_count + sentence_tokens > chunk_size and current_chunk:
                    chunks.append(' '.join(current_chunk))
                    current_chunk = []
                    current_token_count = 0
                
                current_chunk.append(sentence + '.')
                current_token_count += sentence_tokens
        else:
            # If adding this paragraph would exceed chunk size, start a new chunk
            if current_token_count + paragraph_tokens > chunk_size and current_chunk:
                chunks.append(' '.join(current_chunk))
                current_chunk = []
                current_token_count = 0
            
            current_chunk.append(paragraph)
            current_token_count += paragraph_tokens
    
    # Add any remaining text as the last chunk
    if current_chunk:
        chunks.append(' '.join(current_chunk))
    
    print(f"  ✓ Created {len(chunks)} chunks in {time.time() - start_time:.2f} seconds")
    for i, chunk in enumerate(chunks):
        chunk_tokens = len(enc.encode(chunk))
        print(f"    Chunk {i+1}: {chunk_tokens} tokens, {len(chunk)} chars")
    
    return chunks


# Function to generate embeddings
def embed_text(text):
    print(f"  Generating embedding for text ({len(text)} characters)")
    start_time = time.time()
    
    try:
        response = openai_client.embeddings.create(
            model="text-embedding-ada-002", input=text
        )
        embedding = response.data[0].embedding
        vector_dim = len(embedding)
        print(f"  ✓ Embedding generated successfully: {vector_dim} dimensions in {time.time() - start_time:.2f} seconds")
        print(f"  Vector sample (first 5 values): {embedding[:5]}")
        return embedding
    except Exception as e:
        print(f"  ✗ Error generating embedding: {e}")
        raise


# Process and upload emails
print("\n" + "="*80)
print("STARTING DOCUMENT PROCESSING")
print("="*80)

file_count = len([f for f in os.listdir(UNCLEANED_DIR) if f.endswith((".eml", ".txt", ".md"))])
print(f"Found {file_count} .eml, .txt, and .md files to process")

processed_count = 0
success_count = 0
error_count = 0
start_time_all = time.time()

for filename in os.listdir(UNCLEANED_DIR):
    if filename.endswith((".eml", ".txt", ".md")):
        processed_count += 1
        print(f"\n[{processed_count}/{file_count}] Processing: {filename}")
        file_start_time = time.time()
        file_path = os.path.join(UNCLEANED_DIR, filename)
        
        try:
            if filename.endswith(".eml"):
                print(f"Processing as email file")
                cleaned_content, subject, sender, date = clean_email(file_path)
            else:  # For .txt and .md files
                print(f"Processing as {'markdown' if filename.endswith('.md') else 'text'} file")
                try:
                    if filename.endswith(".md"):
                        cleaned_content = extract_markdown_content(file_path)
                    else:
                        with open(file_path, 'r', encoding='utf-8') as file:
                            cleaned_content = file.read()
                    subject = os.path.splitext(filename)[0]  # Use filename as subject
                    sender = "Text File" if filename.endswith(".txt") else "Markdown File"
                    date = os.path.getmtime(file_path)  # Use file modification time as date
                    print(f"  File details:")
                    print(f"    Subject (filename): {subject}")
                    print(f"    Size: {len(cleaned_content)} characters")
                    print(f"    Modified: {datetime.datetime.fromtimestamp(date).strftime('%Y-%m-%d %H:%M:%S')}")
                except Exception as e:
                    print(f"  ✗ Error processing file: {e}")
                    error_count += 1
                    continue

            cleaned_file_path = os.path.join(
                CLEANED_DIR, f"{os.path.splitext(filename)[0]}.txt"
            )
            print(f"  Saving cleaned content to: {cleaned_file_path}")
            with open(cleaned_file_path, "w", encoding="utf-8") as cleaned_file:
                cleaned_file.write(cleaned_content)
            print(f"  ✓ Saved cleaned file ({len(cleaned_content)} characters)")

            # Read the cleaned file and upload its content
            print(f"  Reading cleaned file for upload")
            with open(cleaned_file_path, "r", encoding="utf-8") as cleaned_file:
                text = cleaned_file.read()
            print(f"  ✓ Read {len(text)} characters from cleaned file")

            try:
                print(f"  Generating embeddings and uploading chunks to Pinecone")
                chunks = chunk_text(text)
                
                # Prepare base metadata
                base_metadata = {
                    "subject": subject, 
                    "from": sender, 
                    "date": date,
                    "filename": filename,
                    "processed_at": datetime.datetime.now().isoformat()
                }
                
                # Process and upload each chunk
                upload_start = time.time()
                for i, chunk in enumerate(chunks):
                    chunk_id = f"{filename}_chunk_{i+1}"
                    embedding = embed_text(chunk)
                    
                    # Add chunk-specific metadata
                    chunk_metadata = {
                        **base_metadata,
                        "chunk_index": i + 1,
                        "total_chunks": len(chunks),
                        "chunk_text": chunk
                    }
                    
                    index.upsert([(chunk_id, embedding, chunk_metadata)])
                    print(f"  ✓ Uploaded chunk {i+1}/{len(chunks)}")
                
                print(f"  ✓ Successfully uploaded all chunks to Pinecone in {time.time() - upload_start:.2f} seconds")
                
                success_count += 1
                print(f"  ✓ COMPLETE: Processed {filename} in {time.time() - file_start_time:.2f} seconds")
            except Exception as e:
                print(f"  ✗ ERROR: Failed to process or upload {filename}")
                print(f"  Error details: {str(e)}")
                print(f"  Error type: {type(e).__name__}")
                error_count += 1
        except Exception as e:
            print(f"  ✗ CRITICAL ERROR processing {filename}: {e}")
            error_count += 1

total_time = time.time() - start_time_all
print("\n" + "="*80)
print("PROCESSING SUMMARY")
print("="*80)
print(f"Total files processed: {processed_count}")
print(f"Successfully processed: {success_count}")
print(f"Errors: {error_count}")
print(f"Total processing time: {total_time:.2f} seconds ({total_time/60:.2f} minutes)")
if processed_count > 0:
    print(f"Average time per file: {total_time/processed_count:.2f} seconds")
print(f"Success rate: {success_count/processed_count*100:.1f}% ({success_count}/{processed_count})")
print("="*80)
print(f"SCRIPT COMPLETED: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
print("="*80)
